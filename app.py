import os
import pandas as pd
import streamlit as st
from PIL import Image as PILImage
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage

st.set_page_config(
    page_title="Aplikasi ULP Teknik", page_icon="⚡", layout="wide"
)

EXCEL_FILE = "ULP_TEKNIK.xlsx"
SEGMENT_FILE = "segment.xlsx"
UPLOAD_DIR = "uploads"

# Buat folder uploads jika belum ada
if not os.path.exists(UPLOAD_DIR):
    os.makedirs(UPLOAD_DIR)

def load_data():
    if os.path.exists(EXCEL_FILE):
        xls = pd.ExcelFile(EXCEL_FILE)
        return {sheet: pd.read_excel(EXCEL_FILE, sheet_name=sheet) for sheet in xls.sheet_names}
    return None

def load_master_segment():
    if os.path.exists(SEGMENT_FILE):
        df_seg = pd.read_excel(SEGMENT_FILE)
        df_seg["SECTION"] = df_seg["SECTION"].ffill()
        return df_seg.dropna(subset=["SEGMENT"])
    return None

data_dict = load_data()
df_master_seg = load_master_segment()

st.title("⚡ Sistem Manajemen ULP Teknik (Foto Langsung di Excel)")
menu = st.sidebar.selectbox(
    "Pilih Menu",
    ["Dashboard / Rekap", "Input ROW", "Input Konstruksi", "Input Patroli"]
)

if data_dict:
    if menu == "Dashboard / Rekap":
        st.header("📊 Dashboard & Rekapitulasi Data ULP Teknik")
        tab1, tab2, tab3 = st.tabs(["Data ROW", "Data Konstruksi", "Data Patroli"])
        with tab1:
            st.dataframe(data_dict["ROW"], use_container_width=True)
            st.info(f"Total Data ROW: {len(data_dict['ROW'])}")
        with tab2:
            st.dataframe(data_dict["Konstruksi"], use_container_width=True)
            st.info(f"Total Data Konstruksi: {len(data_dict['Konstruksi'])}")
        with tab3:
            st.dataframe(data_dict["Patroli"], use_container_width=True)
            st.info(f"Total Data Patroli: {len(data_dict['Patroli'])}")

    elif menu == "Input ROW":
        st.header("📝 Input Data ROW (Right of Way)")
        if df_master_seg is not None:
            sec = st.selectbox("Pilih Section", df_master_seg["SECTION"].unique().tolist())
            seg = st.selectbox("Pilih Segment", df_master_seg[df_master_seg["SECTION"] == sec]["SEGMENT"].dropna().tolist())
        else:
            sec, seg = "", st.text_input("Segment")

        with st.form("form_row"):
            tgl = st.date_input("Tanggal")
            tp = st.selectbox("Tebang/Pangkas", ["Tebang", "Pangkas"])
            ket = st.text_area("Keterangan")
            
            f_sebelum = st.file_uploader("Upload Foto Sebelum", type=["jpg", "jpeg", "png"])
            f_sesudah = st.file_uploader("Upload Foto Sesudah", type=["jpg", "jpeg", "png"])
            f_batang = st.file_uploader("Upload Foto Batang Pohon", type=["jpg", "jpeg", "png"])
            
            koord = st.text_input("Koordinat (latitude, longitude)")
            sttus = st.selectbox("Status", ["Pending", "Proses", "Selesai"])

            submitted = st.form_submit_button("Simpan Data ROW & Foto ke Excel")
            
            if submitted:
                path_sebelum = ""
                path_sesudah = ""
                path_batang = ""

                if f_sebelum is not None:
                    path_sebelum = os.path.join(UPLOAD_DIR, f"row_sebelum_{f_sebelum.name}")
                    with open(path_sebelum, "wb") as f:
                        f.write(f_sebelum.getbuffer())

                if f_sesudah is not None:
                    path_sesudah = os.path.join(UPLOAD_DIR, f"row_sesudah_{f_sesudah.name}")
                    with open(path_sesudah, "wb") as f:
                        f.write(f_sesudah.getbuffer())

                if f_batang is not None:
                    path_batang = os.path.join(UPLOAD_DIR, f"row_batang_{f_batang.name}")
                    with open(path_batang, "wb") as f:
                        f.write(f_batang.getbuffer())

                df_row = data_dict["ROW"]
                new_idx = len(df_row) + 2  # baris di excel (header di baris 1)
                
                new_row_data = {
                    "Tanggal": str(tgl),
                    "Section": sec,
                    "Segment": seg,
                    "Tebang/Pangkas": tp,
                    "Keterangan": ket,
                    "Foto Sebelum": path_sebelum,
                    "Foto Sesudah": path_sesudah,
                    "Foto Batang Pohon": path_batang,
                    "Koordinat": koord,
                    "Status": sttus,
                }
                
                updated_df = pd.concat([df_row, pd.DataFrame([new_row_data])], ignore_index=True)
                
                with pd.ExcelWriter(EXCEL_FILE, engine="openpyxl", mode="w") as writer:
                    updated_df.to_excel(writer, sheet_name="ROW", index=False)
                    data_dict["Konstruksi"].to_excel(writer, sheet_name="Konstruksi", index=False)
                    data_dict["Patroli"].to_excel(writer, sheet_name="Patroli", index=False)

                wb = openpyxl.load_workbook(EXCEL_FILE)
                ws = wb["ROW"]

                for col_idx, img_p in [(6, path_sebelum), (7, path_sesudah), (8, path_batang)]:
                    if img_p and os.path.exists(img_p):
                        img_obj = OpenpyxlImage(img_p)
                        img_obj.width = 90
                        img_obj.height = 90
                        cell_coord = f"{openpyxl.utils.get_column_letter(col_idx)}{new_idx}"
                        ws.add_image(img_obj, cell_coord)
                
                ws.row_dimensions[new_idx].height = 75
                wb.save(EXCEL_FILE)

                st.success("Data ROW dan foto fisik berhasil disimpan dan ditampilkan langsung di Excel!")
                st.rerun()

    elif menu == "Input Konstruksi":
        st.header("📝 Input Data Konstruksi")
        with st.form("form_konstruksi"):
            tgl = st.date_input("Tanggal")
            lokasi = st.text_input("Lokasi")
            jenis_peralatan = st.text_input("Jenis Peralatan")
            kondisi = st.selectbox("Kondisi", ["Baik", "Rusak Ringan", "Rusak Berat"])
            
            f_sebelum = st.file_uploader("Upload Foto Sebelum Perbaikan", type=["jpg", "jpeg", "png"])
            f_sesudah = st.file_uploader("Upload Foto Sesudah Perbaikan", type=["jpg", "jpeg", "png"])
            
            status_perbaikan = st.selectbox("Status Perbaikan", ["Belum", "Dalam Perbaikan", "Selesai"])

            submitted = st.form_submit_button("Simpan Konstruksi & Foto")
            if submitted:
                path_sebelum = ""
                path_sesudah = ""
                if f_sebelum is not None:
                    path_sebelum = os.path.join(UPLOAD_DIR, f"konst_sebelum_{f_sebelum.name}")
                    with open(path_sebelum, "wb") as f:
                        f.write(f_sebelum.getbuffer())
                if f_sesudah is not None:
                    path_sesudah = os.path.join(UPLOAD_DIR, f"konst_sesudah_{f_sesudah.name}")
                    with open(path_sesudah, "wb") as f:
                        f.write(f_sesudah.getbuffer())

                df_konst = data_dict["Konstruksi"]
                new_idx = len(df_konst) + 2
                new_row_data = {
                    "Tanggal": str(tgl),
                    "Lokasi": lokasi,
                    "Jenis Peralatan": jenis_peralatan,
                    "Kondisi": kondisi,
                    "Foto Sebelum": path_sebelum,
                    "Foto Sesudah": path_sesudah,
                    "Status Perbaikan": status_perbaikan,
                }
                updated_df = pd.concat([df_konst, pd.DataFrame([new_row_data])], ignore_index=True)

                with pd.ExcelWriter(EXCEL_FILE, engine="openpyxl", mode="w") as writer:
                    data_dict["ROW"].to_excel(writer, sheet_name="ROW", index=False)
                    updated_df.to_excel(writer, sheet_name="Konstruksi", index=False)
                    data_dict["Patroli"].to_excel(writer, sheet_name="Patroli", index=False)

                wb = openpyxl.load_workbook(EXCEL_FILE)
                ws = wb["Konstruksi"]
                for col_idx, img_p in [(5, path_sebelum), (6, path_sesudah)]:
                    if img_p and os.path.exists(img_p):
                        img_obj = OpenpyxlImage(img_p)
                        img_obj.width = 90
                        img_obj.height = 90
                        cell_coord = f"{openpyxl.utils.get_column_letter(col_idx)}{new_idx}"
                        ws.add_image(img_obj, cell_coord)
                ws.row_dimensions[new_idx].height = 75
                wb.save(EXCEL_FILE)

                st.success("Data Konstruksi dan foto berhasil disimpan ke Excel!")
                st.rerun()

    elif menu == "Input Patroli":
        st.header("📝 Input Data Patroli")
        sec_p = st.selectbox("Wilayah / Section Patroli", df_master_seg["SECTION"].unique().tolist()) if df_master_seg is not None else st.text_input("Wilayah Patroli")
        with st.form("form_patroli"):
            tgl = st.date_input("Tanggal")
            petugas = st.text_input("Nama Petugas")
            jumlah_sarpin = st.number_input("Jumlah Sarpin", min_value=0, step=1)
            jumlah_bangtobat = st.number_input("Jumlah Bangtobat", min_value=0, step=1)

            submitted = st.form_submit_button("Simpan Data Patroli")
            if submitted:
                df_patroli = data_dict["Patroli"]
                new_row_data = {
                    "Tanggal": str(tgl),
                    "Petugas": petugas,
                    "Wilayah Patroli": sec_p,
                    "Jumlah Sarpin": jumlah_sarpin,
                    "Jumlah Bangtobat": jumlah_bangtobat,
                }
                updated_df = pd.concat([df_patroli, pd.DataFrame([new_row_data])], ignore_index=True)
                with pd.ExcelWriter(EXCEL_FILE, engine="openpyxl", mode="w") as writer:
                    data_dict["ROW"].to_excel(writer, sheet_name="ROW", index=False)
                    data_dict["Konstruksi"].to_excel(writer, sheet_name="Konstruksi", index=False)
                    updated_df.to_excel(writer, sheet_name="Patroli", index=False)
                st.success("Data Patroli berhasil disimpan!")
                st.rerun()